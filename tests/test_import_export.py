import base64
import csv
import io

import openpyxl

from odoo import Command
from odoo.exceptions import AccessError, UserError
from odoo.tests import TransactionCase, tagged


@tagged("post_install", "-at_install")
class TestErpObeFacultyImportExport(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.Faculty = cls.env["erpobe.master.fakultas"]
        cls.Wizard = cls.env["erpobe.wizard.import.fakultas"]
        cls.admin = cls.env["res.users"].with_context(no_reset_password=True).create(
            {
                "name": "Admin Impor",
                "login": "import.admin@example.test",
                "group_ids": [
                    Command.set([cls.env.ref("erpobe.group_admin").id])
                ],
            }
        )
        cls.superadmin = cls.env["res.users"].with_context(
            no_reset_password=True
        ).create(
            {
                "name": "Superadmin Impor",
                "login": "import.superadmin@example.test",
                "group_ids": [
                    Command.set([cls.env.ref("erpobe.group_superadmin").id])
                ],
            }
        )
        cls.mahasiswa = cls.env["res.users"].with_context(
            no_reset_password=True
        ).create(
            {
                "name": "Mahasiswa Impor",
                "login": "import.mahasiswa@example.test",
                "group_ids": [
                    Command.set([cls.env.ref("erpobe.group_mahasiswa").id])
                ],
            }
        )

    def _new_wizard(self, user=None, **values):
        return self.Wizard.with_user(user or self.superadmin).create(values)

    def _xlsx_upload(self, wizard, values):
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Data"
        worksheet.append(wizard._get_headers())
        worksheet.append(
            [values.get(column["label"], "") for column in wizard._get_columns()]
            + [""]
        )
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        return base64.b64encode(output.getvalue())

    def test_templates_exist_for_all_selected_models(self):
        wizard_models = [
            "erpobe.wizard.import.fakultas",
            "erpobe.wizard.import.prodi",
            "erpobe.wizard.import.dosen",
            "erpobe.wizard.import.mahasiswa",
            "erpobe.wizard.import.rombel",
            "erpobe.wizard.import.matakuliah",
            "erpobe.wizard.import.tahunajaran",
            "erpobe.wizard.import.komponen",
            "erpobe.wizard.import.cpl",
            "erpobe.wizard.import.cpmk",
        ]
        for model_name in wizard_models:
            wizard = self.env[model_name].with_user(self.superadmin).create({})
            wizard.action_download_template()
            content = base64.b64decode(wizard.template_file)
            if wizard.format == "xlsx":
                workbook = openpyxl.load_workbook(
                    io.BytesIO(content), read_only=True
                )
                self.assertIn("Data", workbook.sheetnames, model_name)
                self.assertIn("Panduan", workbook.sheetnames, model_name)
                headers = list(next(workbook["Data"].iter_rows(values_only=True)))
                self.assertEqual(headers, wizard._get_headers(), model_name)
                workbook.close()
            else:
                self.assertTrue(content.startswith(b"\xef\xbb\xbf"), model_name)

    def test_dosen_and_mahasiswa_import_use_profile_create_flow(self):
        faculty = self.Faculty.create(
            {"kode": "FK-USER", "name": "Fakultas User", "active": True}
        )
        prodi = self.env["erpobe.master.prodi"].create(
            {
                "kode": "PR-USER",
                "name": "Prodi User",
                "fakultas_id": faculty.id,
            }
        )
        rombel = self.env["erpobe.master.rombel"].create(
            {"name": "A", "angkatan": 2025, "prodi_id": prodi.id}
        )

        dosen_wizard = self.env["erpobe.wizard.import.dosen"].with_user(
            self.superadmin
        ).create({"dry_run": False})
        dosen_wizard.file = self._xlsx_upload(
            dosen_wizard,
            {
                "NIP": "00001234",
                "Nama Dosen": "Dosen Import",
                "Email": "dosen.import@example.test",
                "Kode Prodi": "PR-USER",
            },
        )
        dosen_wizard.filename = "dosen.xlsx"
        dosen_wizard.action_import()
        dosen = self.env["erpobe.master.dosen"].search(
            [("nip", "=", "00001234")]
        )
        self.assertTrue(dosen)
        self.assertEqual(dosen.prodi_ids, prodi)
        self.assertTrue(dosen.user_id.has_group("erpobe.group_dosen"))
        self.assertEqual(dosen.user_id.login, "dosen.import@example.test")

        mahasiswa_wizard = self.env["erpobe.wizard.import.mahasiswa"].with_user(
            self.superadmin
        ).create({"dry_run": False})
        mahasiswa_wizard.file = self._xlsx_upload(
            mahasiswa_wizard,
            {
                "NIM": "00004567",
                "Nama Mahasiswa": "Mahasiswa Import",
                "Email": "mahasiswa.import@example.test",
                "Kode Prodi": "PR-USER",
                "Angkatan": 2025,
                "Nama Rombel": "A",
            },
        )
        mahasiswa_wizard.filename = "mahasiswa.xlsx"
        mahasiswa_wizard.action_import()
        mahasiswa = self.env["erpobe.master.mahasiswa"].search(
            [("nim", "=", "00004567")]
        )
        self.assertTrue(mahasiswa)
        self.assertEqual(mahasiswa.rombel_id, rombel)
        self.assertTrue(mahasiswa.user_id.has_group("erpobe.group_mahasiswa"))
        self.assertEqual(mahasiswa.user_id.login, "mahasiswa.import@example.test")

    def test_round_trip_for_each_selected_model(self):
        faculty = self.Faculty.create(
            {"kode": "FK-ROUND-ALL", "name": "Fakultas Round Trip", "active": True}
        )
        prodi = self.env["erpobe.master.prodi"].create(
            {
                "kode": "PR-ROUND-ALL",
                "name": "Prodi Round Trip",
                "fakultas_id": faculty.id,
            }
        )
        rombel = self.env["erpobe.master.rombel"].create(
            {
                "name": "A",
                "angkatan": 2025,
                "prodi_id": prodi.id,
            }
        )
        dosen_partner = self.env["res.partner"].create(
            {"name": "Dosen Round Trip", "email": "dosen.round@example.test"}
        )
        dosen = self.env["erpobe.master.dosen"].with_context(
            no_reset_password=True
        ).create(
            {
                "partner_id": dosen_partner.id,
                "nip": "00009876",
                "prodi_ids": [Command.set([prodi.id])],
            }
        )
        mahasiswa_partner = self.env["res.partner"].create(
            {
                "name": "Mahasiswa Round Trip",
                "email": "mahasiswa.round@example.test",
            }
        )
        mahasiswa = self.env["erpobe.master.mahasiswa"].with_context(
            no_reset_password=True
        ).create(
            {
                "partner_id": mahasiswa_partner.id,
                "nim": "00006543",
                "rombel_id": rombel.id,
            }
        )
        matakuliah = self.env["erpobe.master.matakuliah"].create(
            {
                "name": "Mata Kuliah Round Trip",
                "kode_mk": "MK-ROUND",
                "sks": 3,
                "semester": 1,
                "bahasa": "I",
                "prodi_id": prodi.id,
                "koordinator_id": dosen.id,
                "tim_pengampu_ids": [Command.set([dosen.id])],
            }
        )
        cpl = self.env["erpobe.akademik.cpl"].create(
            {
                "name": "Deskripsi CPL Round Trip",
                "kode_cpl": "CPL-ROUND",
                "prodi_id": prodi.id,
            }
        )
        cpmk = self.env["erpobe.akademik.cpmk"].create(
            {
                "mk_id": matakuliah.id,
                "kode_cpmk": "CPMK-ROUND",
                "name": "Deskripsi CPMK Round Trip",
            }
        )
        tahunajaran = self.env["erpobe.akademik.tahunajaran"].create(
            {"tahun_ajaran": "2025/2026", "semester": "ganjil"}
        )
        komponen = self.env["erpobe.akademik.komponen"].create(
            {
                "name": "Komponen Round Trip",
                "tipe": "kognitif",
                "aktif": True,
                "color": 0,
            }
        )

        model_wizards = [
            "erpobe.wizard.import.fakultas",
            "erpobe.wizard.import.prodi",
            "erpobe.wizard.import.dosen",
            "erpobe.wizard.import.mahasiswa",
            "erpobe.wizard.import.rombel",
            "erpobe.wizard.import.matakuliah",
            "erpobe.wizard.import.tahunajaran",
            "erpobe.wizard.import.komponen",
            "erpobe.wizard.import.cpl",
            "erpobe.wizard.import.cpmk",
        ]
        records_by_model = {
            "erpobe.master.fakultas": faculty,
            "erpobe.master.prodi": prodi,
            "erpobe.master.dosen": dosen,
            "erpobe.master.mahasiswa": mahasiswa,
            "erpobe.master.rombel": rombel,
            "erpobe.master.matakuliah": matakuliah,
            "erpobe.akademik.tahunajaran": tahunajaran,
            "erpobe.akademik.komponen": komponen,
            "erpobe.akademik.cpl": cpl,
            "erpobe.akademik.cpmk": cpmk,
        }
        for wizard_model in model_wizards:
            target_model = self.env[wizard_model]._import_export_model
            records = records_by_model[target_model]
            wizard = (
                self.env[wizard_model]
                .with_user(self.superadmin)
                .with_context(
                    active_model=target_model,
                    active_ids=records.ids,
                )
                .create(
                    {
                        "mode": "export",
                        "format": "xlsx",
                        "export_selected": True,
                    }
                )
            )
            wizard.action_export()
            wizard.write(
                {
                    "mode": "import",
                    "file": wizard.export_file,
                    "filename": wizard.export_filename,
                    "dry_run": False,
                }
            )
            wizard.action_import()
            if wizard.error_file:
                report = openpyxl.load_workbook(
                    io.BytesIO(base64.b64decode(wizard.error_file)),
                    read_only=True,
                )
                report_rows = list(report["Kesalahan"].iter_rows(values_only=True))
                report.close()
            else:
                report_rows = []
            self.assertIn(
                "0 dibuat, 1 diperbarui",
                wizard.result_summary,
                f"{wizard_model}: {report_rows}",
            )
            self.assertFalse(wizard.error_file, f"{wizard_model}: {report_rows}")

    def test_template_and_dry_run_do_not_create_sample(self):
        wizard = self._new_wizard(format="xlsx")
        wizard.action_download_template()
        template = base64.b64decode(wizard.template_file)
        workbook = openpyxl.load_workbook(io.BytesIO(template), read_only=True)
        self.assertIn("Data", workbook.sheetnames)
        self.assertIn("Panduan", workbook.sheetnames)
        self.assertIn("Referensi", workbook.sheetnames)
        headers = list(next(workbook["Data"].iter_rows(values_only=True)))
        self.assertEqual(headers, wizard._get_headers())
        self.assertEqual(
            workbook["Data"].cell(row=2, column=len(headers)).value, "CONTOH"
        )
        workbook.close()

        wizard.write(
            {
                "file": wizard.template_file,
                "filename": wizard.template_filename,
                "dry_run": True,
            }
        )
        wizard.action_validate()
        self.assertFalse(self.Faculty.search([("kode", "=", "FK-01")]))
        self.assertIn("0 dibuat", wizard.result_summary)

    def test_csv_import_validates_all_rows_and_reports_original_line(self):
        output = io.StringIO(newline="")
        writer = csv.writer(output, delimiter=";")
        writer.writerow(["Kode Fakultas", "Nama Fakultas", "Aktif", "_contoh"])
        writer.writerow(["FK-CSV", "Fakultas, Utama\nBaris Dua", "Ya", ""])
        writer.writerow(["FK-KOSONG", " ", "Ya", ""])
        upload = base64.b64encode(output.getvalue().encode("utf-8-sig"))

        wizard = self._new_wizard(
            format="csv",
            file=upload,
            filename="fakultas.csv",
            dry_run=True,
        )
        wizard.action_validate()
        self.assertFalse(self.Faculty.search([("kode", "=", "FK-CSV")]))
        self.assertIn("1 dibuat", wizard.result_summary)
        self.assertIn("1 ditolak", wizard.result_summary)

        wizard.dry_run = False
        wizard.action_import()
        faculty = self.Faculty.search([("kode", "=", "FK-CSV")])
        self.assertEqual(faculty.name, "Fakultas, Utama\nBaris Dua")
        report = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(wizard.error_file)), read_only=True
        )
        error_row = list(report["Kesalahan"].iter_rows(values_only=True))[1]
        self.assertEqual(error_row[0], 4)
        self.assertIn("Nama Fakultas wajib diisi", error_row[2])
        report.close()

    def test_import_rejects_duplicate_keys_missing_headers_and_unknown_relations(self):
        wizard = self._new_wizard(format="xlsx", dry_run=True)
        duplicate_row = {
            "Kode Fakultas": "FK-DUP-IMPORT",
            "Nama Fakultas": "Fakultas Duplikat",
            "Aktif": "Ya",
        }
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Data"
        worksheet.append(wizard._get_headers())
        worksheet.append(
            [duplicate_row.get(column["label"], "") for column in wizard._get_columns()]
            + [""]
        )
        worksheet.append(
            [duplicate_row.get(column["label"], "") for column in wizard._get_columns()]
            + [""]
        )
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        wizard.file = base64.b64encode(output.getvalue())
        wizard.filename = "duplicate.xlsx"
        wizard.action_validate()
        self.assertIn("2 ditolak", wizard.result_summary)
        self.assertFalse(self.Faculty.search([("kode", "=", "FK-DUP-IMPORT")]))

        missing_header = self._new_wizard(
            format="csv",
            file=base64.b64encode(
                b"Kode Fakultas,Aktif\nFK-MISSING,Ya\n"
            ),
            filename="missing.csv",
        )
        with self.assertRaisesRegex(UserError, "Kolom wajib tidak ditemukan"):
            missing_header.action_validate()

        unknown_relation = self.env["erpobe.wizard.import.prodi"].with_user(
            self.superadmin
        ).create(
            {
                "format": "xlsx",
                "dry_run": True,
            }
        )
        unknown_relation.file = self._xlsx_upload(
            unknown_relation,
            {
                "Kode Fakultas": "FK-NOT-FOUND-TRITON-2026",
                "Kode Prodi": "PR-NOT-FOUND-TRITON-2026",
                "Nama Prodi": "Program Studi Impor",
                "Aktif": "Ya",
            },
        )
        unknown_relation.filename = "unknown_relation.xlsx"
        unknown_relation.action_validate()
        self.assertIn("1 ditolak", unknown_relation.result_summary)

    def test_import_rejects_mismatched_format_and_oversized_file(self):
        format_mismatch = self._new_wizard(
            format="xlsx",
            file=base64.b64encode(b"Kode Fakultas,Nama Fakultas\n"),
            filename="fakultas.csv",
        )
        with self.assertRaisesRegex(UserError, "Format file tidak sesuai"):
            format_mismatch.action_validate()

        oversized = self._new_wizard(
            format="csv",
            file=base64.b64encode(b"x" * (5 * 1024 * 1024 + 1)),
            filename="terlalu_besar.csv",
        )
        with self.assertRaisesRegex(UserError, "Ukuran file melebihi batas"):
            oversized.action_validate()

    def test_export_round_trip_is_idempotent(self):
        faculty = self.Faculty.create(
            {"kode": "FK-ROUND", "name": "Fakultas Round Trip", "active": True}
        )
        wizard = self._new_wizard(
            mode="export",
            format="xlsx",
            export_selected=True,
        )
        wizard = wizard.with_context(
            active_model="erpobe.master.fakultas",
            active_ids=[faculty.id],
        )
        wizard.action_export()
        export_file = wizard.export_file
        export_filename = wizard.export_filename
        workbook = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(export_file)), read_only=True
        )
        self.assertEqual(
            list(next(workbook["Data"].iter_rows(values_only=True))),
            wizard._get_headers(),
        )
        workbook.close()

        wizard.write(
            {
                "mode": "import",
                "file": export_file,
                "filename": export_filename,
                "format": "xlsx",
                "dry_run": False,
                "update_existing": True,
            }
        )
        wizard.action_import()
        self.assertEqual(
            self.Faculty.search_count([("kode", "=", "FK-ROUND")]), 1
        )
        self.assertIn("0 dibuat, 1 diperbarui", wizard.result_summary)
        self.assertEqual(faculty.name, "Fakultas Round Trip")

    def test_non_admin_cannot_run_wizard(self):
        wizard = self.Wizard.with_user(self.admin).create({})
        with self.assertRaises(AccessError):
            wizard.with_user(self.mahasiswa).action_export()

    def test_admin_can_download_template(self):
        wizard = self._new_wizard(user=self.admin, format="csv")
        wizard.action_download_template()
        self.assertTrue(wizard.template_file)

    def test_unknown_headers_and_formula_like_text(self):
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(["Kode Fakultas", "Nama Fakultas", "Aktif", "Kolom Aneh"])
        writer.writerow(["FK-UNKNOWN", "=1+1", "Ya", "nilai"])
        wizard = self._new_wizard(
            format="csv",
            file=base64.b64encode(output.getvalue().encode("utf-8")),
            filename="unknown.csv",
        )
        with self.assertRaisesRegex(UserError, "Kolom tidak dikenal"):
            wizard.action_validate()
        self.assertFalse(self.Faculty.search([("kode", "=", "FK-UNKNOWN")]))

        faculty = self.Faculty.create(
            {"kode": "FK-FORMULA", "name": "=1+1", "active": True}
        )
        export_wizard = self._new_wizard(mode="export", format="xlsx")
        export_wizard.action_export()
        workbook = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(export_wizard.export_file)), data_only=False
        )
        values = list(workbook["Data"].iter_rows(values_only=True))
        name_column = export_wizard._get_headers().index("Nama Fakultas")
        exported = next(
            row[name_column] for row in values[1:] if row[0] == faculty.kode
        )
        self.assertEqual(exported, "'=1+1")
        formula_row = next(
            row_index
            for row_index, row in enumerate(values[1:], start=2)
            if row[0] == faculty.kode
        )
        self.assertEqual(
            workbook["Data"].cell(row=formula_row, column=name_column + 1).data_type,
            "s",
        )
        workbook.close()
        import_wizard = self._new_wizard(
            mode="import",
            format="xlsx",
            file=export_wizard.export_file,
            filename=export_wizard.export_filename,
            dry_run=False,
        )
        import_wizard.action_import()
        self.assertEqual(faculty.name, "=1+1")
