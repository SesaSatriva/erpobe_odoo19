import base64
import csv
import html
import io
import re
import zipfile
from datetime import datetime
from urllib.parse import quote

import openpyxl
import xlsxwriter

from odoo import api, fields, models
from odoo.exceptions import AccessError, UserError


MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5000


class _DryRunRollback(Exception):
    pass


class ErpObeImportExportMixin(models.AbstractModel):
    _name = "erpobe.import.export.mixin"
    _description = "Mesin Import dan Export ERP OBE"

    mode = fields.Selection(
        [("import", "Impor"), ("export", "Ekspor")],
        default="import",
        required=True,
    )
    file = fields.Binary(string="File", attachment=False)
    filename = fields.Char(string="Nama File")
    format = fields.Selection(
        [("xlsx", "Excel (.xlsx)"), ("csv", "CSV (.csv)")],
        default="xlsx",
        required=True,
    )
    dry_run = fields.Boolean(string="Validasi saja", default=True)
    update_existing = fields.Boolean(
        string="Perbarui data yang sudah ada", default=True
    )
    stop_on_error = fields.Boolean(string="Batalkan semua jika ada kesalahan")
    export_selected = fields.Boolean(string="Ekspor record terpilih")
    filter_fakultas_id = fields.Many2one(
        "erpobe.master.fakultas", string="Filter Fakultas"
    )
    filter_prodi_id = fields.Many2one(
        "erpobe.master.prodi", string="Filter Program Studi"
    )
    filter_fakultas_available = fields.Boolean(
        compute="_compute_filter_availability"
    )
    filter_prodi_available = fields.Boolean(compute="_compute_filter_availability")
    state = fields.Selection(
        [
            ("draft", "Draft"),
            ("validated", "Tervalidasi"),
            ("done", "Selesai"),
        ],
        default="draft",
        readonly=True,
    )
    result_summary = fields.Html(string="Ringkasan", readonly=True)
    error_file = fields.Binary(string="Laporan Kesalahan", readonly=True)
    error_filename = fields.Char(string="Nama Laporan", readonly=True)
    template_file = fields.Binary(string="Template", readonly=True)
    template_filename = fields.Char(string="Nama Template", readonly=True)
    export_file = fields.Binary(string="Hasil Ekspor", readonly=True)
    export_filename = fields.Char(string="Nama Hasil Ekspor", readonly=True)
    guide = fields.Html(compute="_compute_guide", string="Panduan")

    _import_export_model = None
    _import_export_columns = ()
    _import_export_identity = ()

    @api.depends()
    def _compute_guide(self):
        for wizard in self:
            rows = []
            for column in wizard._get_columns():
                requirement = "Wajib" if column.get("required") else "Opsional"
                rows.append(
                    "<tr>"
                    f"<td>{html.escape(column['label'])}</td>"
                    f"<td>{html.escape(column['type'])}</td>"
                    f"<td>{requirement}</td>"
                    f"<td>{html.escape(str(column.get('example', '')))}</td>"
                    f"<td>{html.escape(column.get('description', ''))}</td>"
                    "</tr>"
                )
            wizard.guide = (
                "<p>Unduh template untuk melihat contoh dan struktur kolom. "
                "Untuk CSV, baris panduan yang diawali # akan diabaikan saat impor. "
                "Nilai relasi menggunakan kode atau identitas yang dijelaskan.</p>"
                "<table class='table table-sm table-striped'><thead><tr>"
                "<th>Kolom</th><th>Tipe</th><th>Kewajiban</th>"
                "<th>Contoh</th><th>Keterangan</th>"
                "</tr></thead><tbody>"
                + "".join(rows)
                + "</tbody></table>"
            )

    @api.depends()
    def _compute_filter_availability(self):
        for wizard in self:
            model = wizard._get_target_model()
            wizard.filter_fakultas_available = (
                model._name == "erpobe.master.fakultas"
                or "fakultas_id" in model._fields
                or "prodi_id" in model._fields
                or "prodi_ids" in model._fields
            )
            wizard.filter_prodi_available = (
                "prodi_id" in model._fields or "prodi_ids" in model._fields
            )

    def _check_wizard_access(self):
        if not (
            self.env.user.has_group("erpobe.group_admin")
            or self.env.user.has_group("erpobe.group_superadmin")
        ):
            raise AccessError(
                "Hanya Admin atau Superadmin yang dapat menggunakan wizard ini."
            )

    def _get_columns(self):
        return self._import_export_columns

    def _get_target_model(self):
        if not self._import_export_model:
            raise UserError("Model untuk impor atau ekspor belum ditentukan.")
        return self.env[self._import_export_model]

    def _get_headers(self):
        return [column["label"] for column in self._get_columns()] + ["_contoh"]

    @staticmethod
    def _normalize_header(value):
        return re.sub(r"[\s_]+", "", str(value or "").strip().lower())

    @staticmethod
    def _is_blank(value):
        return value is None or (isinstance(value, str) and not value.strip())

    def _decode_upload(self):
        self.ensure_one()
        if not self.file or not self.filename:
            raise UserError("Pilih file .xlsx atau .csv untuk diimpor.")
        if len(self.file) > ((MAX_FILE_SIZE + 2) // 3) * 4:
            raise UserError("Ukuran file melebihi batas 5 MB.")
        try:
            raw = base64.b64decode(self.file, validate=True)
        except (ValueError, TypeError):
            raise UserError("Isi file tidak valid.") from None
        if len(raw) > MAX_FILE_SIZE:
            raise UserError("Ukuran file melebihi batas 5 MB.")

        extension = self.filename.lower().rsplit(".", 1)[-1]
        if extension != self.format:
            raise UserError("Format file tidak sesuai dengan pilihan format.")
        if extension == "xlsx":
            try:
                with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                    members = set(archive.namelist())
                    if (
                        "[Content_Types].xml" not in members
                        or "xl/workbook.xml" not in members
                        or "xl/vbaProject.bin" in members
                    ):
                        raise UserError("File bukan dokumen Excel .xlsx yang valid.")
            except (zipfile.BadZipFile, OSError):
                raise UserError("File Excel rusak atau tidak valid.") from None
            return self._read_xlsx(raw)
        if extension == "csv":
            return self._read_csv(raw)
        raise UserError("Format file tidak didukung. Gunakan .xlsx atau .csv.")

    def _read_xlsx(self, raw):
        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(raw), read_only=True, data_only=False
            )
            if "Data" not in workbook.sheetnames:
                raise UserError("File Excel harus memiliki sheet bernama Data.")
            worksheet = workbook["Data"]
            rows = worksheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                raise UserError("Sheet Data tidak memiliki header.")
            data_rows = [
                (row_number, row)
                for row_number, row in enumerate(rows, start=2)
                if any(not self._is_blank(value) for value in row)
            ]
            workbook.close()
            return headers, data_rows
        except (OSError, ValueError, KeyError, zipfile.BadZipFile):
            raise UserError("File Excel rusak atau tidak dapat dibaca.") from None

    def _read_csv(self, raw):
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = raw.decode("cp1252")
            except UnicodeDecodeError:
                raise UserError(
                    "CSV harus menggunakan UTF-8 atau Windows-1252."
                ) from None
        physical_lines = text.splitlines(keepends=True)
        skipped_lines = 0
        while (
            skipped_lines < len(physical_lines)
            and physical_lines[skipped_lines].lstrip().startswith("#")
        ):
            skipped_lines += 1
        content = "".join(physical_lines[skipped_lines:])
        if not content.strip():
            raise UserError("File CSV tidak memiliki header.")
        sample = "\n".join(content.splitlines()[:10])
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;")
        except csv.Error:
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
            reader = csv.reader(
                io.StringIO(content, newline=""), delimiter=delimiter
            )
        else:
            reader = csv.reader(io.StringIO(content, newline=""), dialect)
        headers = next(reader, None)
        if not headers:
            raise UserError("File CSV tidak memiliki header.")
        data_rows = []
        while True:
            previous_line = reader.line_num
            try:
                row = next(reader)
            except StopIteration:
                break
            if any(not self._is_blank(value) for value in row):
                data_rows.append(
                    (skipped_lines + previous_line + 1, row)
                )
        return headers, data_rows

    def _prepare_headers(self, headers):
        known = {
            self._normalize_header(column["label"]): column
            for column in self._get_columns()
        }
        known[self._normalize_header("_contoh")] = None
        mapped = {}
        unknown = []
        duplicates = []
        for position, header in enumerate(headers):
            normalized = self._normalize_header(header)
            if normalized in mapped:
                duplicates.append(str(header))
            elif normalized not in known:
                unknown.append(str(header))
            else:
                mapped[normalized] = position
        missing = [
            column["label"]
            for column in self._get_columns()
            if column.get("required")
            and self._normalize_header(column["label"]) not in mapped
        ]
        messages = []
        if missing:
            messages.append("Kolom wajib tidak ditemukan: " + ", ".join(missing))
        if unknown:
            messages.append("Kolom tidak dikenal: " + ", ".join(unknown))
        if duplicates:
            messages.append("Header duplikat: " + ", ".join(duplicates))
        if messages:
            raise UserError(". ".join(messages) + ".")
        return mapped

    def _is_example(self, row, header_map):
        index = header_map.get(self._normalize_header("_contoh"))
        if index is None or index >= len(row):
            return False
        value = str(row[index] or "").strip().lower()
        return value in {"1", "true", "ya", "contoh", "sample"}

    def _convert_value(self, column, value, relation_cache):
        if self._is_blank(value):
            if column.get("required"):
                raise ValueError(f"{column['label']} wajib diisi.")
            return None

        value_type = column["type"]
        if value_type in {"char", "text"}:
            text = str(value).strip()
            if text.startswith("'") and text[1:].startswith(
                ("=", "+", "-", "@", "\t", "\r")
            ):
                return text[1:]
            return text
        if value_type == "integer":
            try:
                number = int(value)
            except (TypeError, ValueError):
                raise ValueError(
                    f"{column['label']} harus berupa bilangan bulat."
                ) from None
            return number
        if value_type == "boolean":
            if isinstance(value, bool):
                return value
            normalized = str(value).strip().lower()
            if normalized in {"1", "true", "ya", "yes", "aktif"}:
                return True
            if normalized in {"0", "false", "tidak", "no", "nonaktif"}:
                return False
            raise ValueError(f"{column['label']} harus Ya/Tidak atau True/False.")
        if value_type == "selection":
            selections = dict(column["selection"])
            normalized = str(value).strip().lower()
            for key, label in selections.items():
                if normalized in {str(key).lower(), str(label).lower()}:
                    return key
            allowed = ", ".join(selections)
            raise ValueError(f"{column['label']} harus salah satu: {allowed}.")
        if value_type == "many2one":
            return self._resolve_relation(column, value, relation_cache)
        if value_type == "many2many":
            values = [part.strip() for part in str(value).split(";") if part.strip()]
            return [
                self._resolve_relation(column, item, relation_cache) for item in values
            ]
        raise UserError(f"Tipe kolom {value_type} belum didukung.")

    def _resolve_relation(self, column, value, relation_cache):
        model_name = column["relation"]
        lookup_field = column.get("lookup_field", "name")
        key = (model_name, lookup_field, str(value).strip())
        if key in relation_cache:
            return relation_cache[key]
        model = self.env[model_name]
        matches = model.search([(lookup_field, "=", str(value).strip())], limit=2)
        if not matches:
            raise ValueError(
                f"{column['label']}: {value!s} tidak ditemukan pada "
                f"{model._description}."
            )
        if len(matches) > 1:
            raise ValueError(
                f"{column['label']}: {value!s} cocok dengan lebih dari satu data."
            )
        relation_cache[key] = matches.id
        return matches.id

    def _prepare_import_values(self, row_values):
        values = {}
        for column in self._get_columns():
            value = row_values[column["label"]]
            if value is None and not column.get("required"):
                continue
            field_name = column.get("field", column["name"])
            if column["type"] == "many2many":
                values[field_name] = [(6, 0, value)]
            else:
                values[field_name] = value
        return values

    def _create_record(self, values, dry_run=False):
        target = self._get_target_model()
        if dry_run:
            target = target.with_context(no_reset_password=True)
        return target.create(values)

    def _write_record(self, record, values, dry_run=False):
        if dry_run:
            record = record.with_context(no_reset_password=True)
        return record.write(values)

    def _find_existing(self, values):
        domain = []
        for field_name in self._import_export_identity:
            value = values.get(field_name)
            if value is None:
                return self._get_target_model().browse()
            domain.append((field_name, "=", value))
        if not domain:
            return self._get_target_model().browse()
        return self._get_target_model().search(domain, limit=2)

    def _natural_key(self, values):
        return tuple(
            values.get(field_name) for field_name in self._import_export_identity
        )

    def _action_import(self):
        self.ensure_one()
        self._check_wizard_access()
        target = self._get_target_model()
        headers, rows = self._decode_upload()
        if len(rows) > MAX_IMPORT_ROWS:
            raise UserError(f"File melebihi batas {MAX_IMPORT_ROWS} baris data.")
        header_map = self._prepare_headers(headers)
        column_positions = {
            column["label"]: header_map[self._normalize_header(column["label"])]
            for column in self._get_columns()
        }

        prepared = []
        errors = []
        seen_keys = {}
        relation_cache = {}
        for line_number, row in rows:
            if self._is_example(row, header_map):
                continue
            row_values = {}
            row_errors = []
            if len(row) > len(headers) and any(
                not self._is_blank(value) for value in row[len(headers) :]
            ):
                row_errors.append(
                    ("Baris", "Jumlah kolom data melebihi jumlah header.")
                )
            for column in self._get_columns():
                position = column_positions[column["label"]]
                value = row[position] if position < len(row) else None
                try:
                    row_values[column["label"]] = self._convert_value(
                        column, value, relation_cache
                    )
                except ValueError as error:
                    row_errors.append((column["label"], str(error)))
            if row_errors:
                errors.extend(
                    {"row": line_number, "column": column, "message": message}
                    for column, message in row_errors
                )
                continue
            try:
                values = self._prepare_import_values(row_values)
            except ValueError as error:
                errors.append(
                    {
                        "row": line_number,
                        "column": "Relasi",
                        "message": str(error),
                    }
                )
                continue
            key = self._natural_key(values)
            previous = seen_keys.get(key)
            if previous:
                prior_item = previous.get("prepared")
                if prior_item in prepared:
                    prepared.remove(prior_item)
                for duplicate_line in (previous["row"], line_number):
                    errors.append(
                        {
                            "row": duplicate_line,
                            "column": ", ".join(self._import_export_identity),
                            "message": "Kunci alami duplikat dalam file.",
                        }
                    )
                continue
            seen_keys[key] = {"row": line_number, "prepared": None}
            existing = self._find_existing(values)
            if len(existing) > 1:
                errors.append(
                    {
                        "row": line_number,
                        "column": ", ".join(self._import_export_identity),
                        "message": "Kunci alami cocok dengan lebih dari satu record.",
                    }
                )
                continue
            if existing and not self.update_existing:
                item = {
                    "row": line_number,
                    "values": values,
                    "existing": existing,
                    "skip": True,
                }
            else:
                item = {
                    "row": line_number,
                    "values": values,
                    "existing": existing,
                    "skip": False,
                }
            prepared.append(item)
            identity_key = self._natural_key(values)
            seen_keys[identity_key]["prepared"] = item

        created = updated = skipped = 0
        if self.dry_run:
            for item in prepared:
                if item["skip"]:
                    skipped += 1
                    continue
                try:
                    with self.env.cr.savepoint():
                        if item["existing"]:
                            self._write_record(
                                item["existing"], item["values"], dry_run=True
                            )
                        else:
                            self._create_record(item["values"], dry_run=True)
                        raise _DryRunRollback()
                except _DryRunRollback:
                    if item["existing"]:
                        updated += 1
                    else:
                        created += 1
                except Exception as error:
                    errors.append(self._format_row_error(item["row"], error))
        else:
            try:
                with self.env.cr.savepoint():
                    for item in prepared:
                        if item["skip"]:
                            skipped += 1
                            continue
                        try:
                            with self.env.cr.savepoint():
                                if item["existing"]:
                                    self._write_record(
                                        item["existing"], item["values"]
                                    )
                                    updated += 1
                                else:
                                    self._create_record(item["values"])
                                    created += 1
                        except Exception as error:
                            errors.append(self._format_row_error(item["row"], error))
                            if self.stop_on_error:
                                raise _DryRunRollback()
                    if self.stop_on_error and errors:
                        raise _DryRunRollback()
            except _DryRunRollback:
                created = updated = 0
                skipped = 0

        self.error_file = False
        self.error_filename = False
        if errors:
            self.error_file = base64.b64encode(self._make_error_report(errors))
            self.error_filename = f"laporan_kesalahan_{self._safe_model_name()}.xlsx"
        self.state = "validated" if self.dry_run else "done"
        summary = (
            f"{'Validasi selesai' if self.dry_run else 'Impor selesai'}: "
            f"{created} dibuat, {updated} diperbarui, {skipped} dilewati, "
            f"{len(errors)} ditolak."
        )
        if self.dry_run:
            summary += " Tidak ada perubahan data karena mode validasi."
        self.result_summary = f"<p>{summary}</p>"
        return self._reload_action()

    @staticmethod
    def _format_row_error(row_number, error):
        if isinstance(error, (AccessError, UserError)):
            message = str(error)
        else:
            message = str(error).split("\n", 1)[0]
        return {"row": row_number, "column": "Baris", "message": message}

    def _make_error_report(self, errors):
        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {"in_memory": True})
        worksheet = workbook.add_worksheet("Kesalahan")
        headers = ["Nomor Baris", "Kolom", "Pesan"]
        for col, header in enumerate(headers):
            worksheet.write_string(0, col, header)
        for row_number, error in enumerate(errors, start=1):
            worksheet.write_number(row_number, 0, error["row"])
            worksheet.write_string(row_number, 1, str(error["column"]))
            worksheet.write_string(
                row_number, 2, self._escape_formula(error["message"])
            )
        workbook.close()
        return output.getvalue()

    def action_validate(self):
        self.ensure_one()
        self.dry_run = True
        return self._action_import()

    def action_import(self):
        self.ensure_one()
        return self._action_import()

    def action_download_template(self):
        self.ensure_one()
        self._check_wizard_access()
        self.template_file = base64.b64encode(self._make_template())
        self.template_filename = f"template_{self._safe_model_name()}.{self.format}"
        return self._download_action("template_file", "template_filename")

    def _make_template(self):
        headers = self._get_headers()
        if self.format == "csv":
            output = io.StringIO(newline="")
            output.write(
                "# Panduan: isi kolom sesuai header; hapus atau pertahankan "
                "baris CONTOH, baris itu tidak akan diimpor.\r\n"
            )
            output.write(
                "# Nilai relasi menggunakan identitas yang tercantum di wizard.\r\n"
            )
            writer = csv.writer(output)
            writer.writerow(headers)
            example = [
                column.get("example", "") for column in self._get_columns()
            ] + ["CONTOH"]
            writer.writerow(example)
            return output.getvalue().encode("utf-8-sig")

        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {"in_memory": True})
        data_sheet = workbook.add_worksheet("Data")
        guide_sheet = workbook.add_worksheet("Panduan")
        header_format = workbook.add_format(
            {"bold": True, "bg_color": "#D9EAD3", "border": 1}
        )
        example_format = workbook.add_format(
            {"bg_color": "#FFF2CC", "font_color": "#7F6000"}
        )
        for col, header in enumerate(headers):
            data_sheet.write_string(0, col, header, header_format)
            data_sheet.set_column(col, col, max(16, min(len(header) + 4, 36)))
        for col, spec in enumerate(self._get_columns()):
            value = spec.get("example", "")
            data_sheet.write_string(1, col, str(value), example_format)
        data_sheet.write_string(
            1, len(headers) - 1, "CONTOH", example_format
        )
        text_format = workbook.add_format({"num_format": "@"})
        for col, spec in enumerate(self._get_columns()):
            if spec["type"] in {"char", "text", "many2one", "many2many"}:
                data_sheet.set_column(
                    col, col, max(16, min(len(spec["label"]) + 4, 36)), text_format
                )
            elif spec["type"] == "boolean":
                data_sheet.data_validation(
                    1,
                    col,
                    MAX_IMPORT_ROWS,
                    col,
                    {"validate": "list", "source": ["Ya", "Tidak"]},
                )
            elif spec["type"] == "selection":
                data_sheet.data_validation(
                    1,
                    col,
                    MAX_IMPORT_ROWS,
                    col,
                    {
                        "validate": "list",
                        "source": [label for _, label in spec["selection"]],
                    },
                )
        data_sheet.freeze_panes(1, 0)
        guide_headers = [
            "Kolom",
            "Tipe",
            "Wajib",
            "Contoh",
            "Keterangan",
        ]
        guide_sheet.write_row(0, 0, guide_headers, header_format)
        for row, spec in enumerate(self._get_columns(), start=1):
            guide_sheet.write_row(
                row,
                0,
                [
                    spec["label"],
                    spec["type"],
                    "Ya" if spec.get("required") else "Tidak",
                    spec.get("example", ""),
                    spec.get("description", ""),
                ],
            )
        guide_sheet.set_column(0, 0, 24)
        guide_sheet.set_column(1, 3, 18)
        guide_sheet.set_column(4, 4, 64)
        relation_columns = [
            column
            for column in self._get_columns()
            if column["type"] in {"many2one", "many2many"}
        ]
        reference_sheet = workbook.add_worksheet("Referensi")
        if relation_columns:
            for column_index, spec in enumerate(relation_columns):
                relation_model = self.env[spec["relation"]]
                lookup_field = spec.get("lookup_field", "name")
                records = relation_model.with_context(active_test=False).search(
                    [], order=lookup_field, limit=MAX_IMPORT_ROWS
                )
                reference_sheet.write_string(
                    0, column_index * 2, spec["label"], header_format
                )
                reference_sheet.write_string(
                    0, column_index * 2 + 1, "Nama", header_format
                )
                for row_index, record in enumerate(records, start=1):
                    reference_sheet.write_string(
                        row_index, column_index * 2, str(record[lookup_field] or "")
                    )
                    reference_sheet.write_string(
                        row_index,
                        column_index * 2 + 1,
                        str(record.display_name or ""),
                    )
                reference_sheet.set_column(column_index * 2, column_index * 2 + 1, 28)
        else:
            reference_sheet.write_string(
                0, 0, "Tidak ada kolom relasi untuk referensi."
            )
        workbook.close()
        return output.getvalue()

    def action_export(self):
        self.ensure_one()
        self._check_wizard_access()
        model = self._get_target_model()
        domain = []
        active_ids = self.env.context.get("active_ids")
        active_model = self.env.context.get("active_model")
        if (
            self.export_selected
            and active_ids
            and active_model == self._import_export_model
        ):
            domain.append(("id", "in", active_ids))
        if self.filter_fakultas_id:
            if model._name == "erpobe.master.fakultas":
                domain.append(("id", "=", self.filter_fakultas_id.id))
            elif "fakultas_id" in model._fields:
                domain.append(("fakultas_id", "=", self.filter_fakultas_id.id))
            elif "prodi_id" in model._fields:
                domain.append(
                    ("prodi_id.fakultas_id", "=", self.filter_fakultas_id.id)
                )
            elif "prodi_ids" in model._fields:
                domain.append(
                    ("prodi_ids.fakultas_id", "=", self.filter_fakultas_id.id)
                )
            else:
                raise UserError("Model ini tidak mendukung filter fakultas.")
        if self.filter_prodi_id:
            if "prodi_id" in model._fields:
                domain.append(("prodi_id", "=", self.filter_prodi_id.id))
            elif "prodi_ids" in model._fields:
                domain.append(("prodi_ids", "in", self.filter_prodi_id.id))
            else:
                raise UserError("Model ini tidak mendukung filter program studi.")
        records = model.with_context(active_test=False).search(domain, order="id")
        if len(records) > MAX_IMPORT_ROWS:
            raise UserError(f"Ekspor dibatasi hingga {MAX_IMPORT_ROWS} record.")
        rows = [self._export_record(record) for record in records]
        self.export_file = base64.b64encode(self._write_export(rows))
        timestamp = datetime.now().strftime("%Y%m%d_%H%M")
        self.export_filename = (
            f"export_{self._safe_model_name()}_{timestamp}.{self.format}"
        )
        return self._download_action("export_file", "export_filename")

    def action_download_error_report(self):
        self.ensure_one()
        self._check_wizard_access()
        if not self.error_file:
            raise UserError("Belum ada laporan kesalahan untuk diunduh.")
        return self._download_action("error_file", "error_filename")

    def _export_record(self, record):
        values = []
        for column in self._get_columns():
            field_name = column.get("field", column["name"])
            field = record._fields.get(field_name)
            value = record
            for path_part in column.get("export_field", field_name).split("."):
                value = value[path_part]
            if column["type"] == "many2one":
                value = value[column.get("lookup_field", "name")] if value else ""
            elif column["type"] == "many2many":
                value = ";".join(
                    str(item[column.get("lookup_field", "name")])
                    for item in value
                )
            elif column["type"] == "boolean":
                value = "Ya" if value else "Tidak"
            elif column["type"] == "selection":
                value = (
                    dict(field.selection).get(value, "")
                    if field and not value
                    else value
                )
            values.append(value if value is not None else "")
        values.append("")
        return values

    def _write_export(self, rows):
        headers = self._get_headers()
        if self.format == "csv":
            output = io.StringIO(newline="")
            writer = csv.writer(output)
            writer.writerow(headers)
            for row in rows:
                writer.writerow(
                    [
                        self._escape_formula(value)
                        if isinstance(value, str)
                        else value
                        for value in row
                    ]
                )
            return output.getvalue().encode("utf-8-sig")
        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {"in_memory": True})
        worksheet = workbook.add_worksheet("Data")
        for col, header in enumerate(headers):
            worksheet.write_string(0, col, header)
        for row_index, row in enumerate(rows, start=1):
            for col_index, value in enumerate(row):
                if isinstance(value, str):
                    worksheet.write_string(
                        row_index, col_index, self._escape_formula(value)
                    )
                elif isinstance(value, bool):
                    worksheet.write_boolean(row_index, col_index, value)
                elif isinstance(value, (int, float)):
                    worksheet.write_number(row_index, col_index, value)
                else:
                    worksheet.write_string(
                        row_index, col_index, str(value) if value else ""
                    )
        workbook.close()
        return output.getvalue()

    @staticmethod
    def _escape_formula(value):
        text = str(value)
        if text.startswith(("=", "+", "-", "@", "\t", "\r")):
            return "'" + text
        return text

    def _safe_model_name(self):
        return re.sub(r"[^a-z0-9_]+", "_", self._import_export_model.split(".")[-1])

    def _download_action(self, field_name, filename_field):
        url = (
            "/web/content?"
            f"model={quote(self._name)}&id={self.id}"
            f"&field={field_name}&filename_field={filename_field}&download=true"
        )
        return {"type": "ir.actions.act_url", "url": url, "target": "self"}

    @staticmethod
    def _reload_action():
        return {"type": "ir.actions.client", "tag": "reload"}
